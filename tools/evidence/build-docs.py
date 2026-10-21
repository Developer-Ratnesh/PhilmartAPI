"""Writes the Milestone 2 developer response and evidence submission, as DOCX
(for editing) and HTML (printed to PDF by print-pdf.mjs), from one source.

  python -I tools/evidence/build-docs.py --facts 00_INDEX.facts.json --out <folder>
      --release m2-rc2 --commit <sha> --web <url> --api <url> --date "22 October 2026"
"""

import argparse
import html
import json
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from catalogue import ATTACHMENT_A, ATTACHMENT_B  # noqa: E402

NAVY = "16233A"
BLUE = "2F6DA8"


# --------------------------------------------------------------- content

def content(f, a):
    """The document as a list of blocks: (kind, payload)."""
    counts = f["counts"]
    rows = f["rows"]
    auto = f["automated"]
    runs = f["runs"]
    live = next((r for r in runs if r["label"] == "Live"), None)
    env = live or (runs[0] if runs else None)
    env_label = "the live staging site" if live else "a local copy of the full stack (pre-deployment verification)"

    def c(src, k):
        return counts.get(src, {}).get(k, 0)

    def total(src):
        return sum(counts.get(src, {}).values())

    def outcome(tid):
        return rows.get(tid, {}).get("outcome", "NOT RUN")

    B = []
    add = B.append

    add(("title", "Milestone 2: developer response and evidence submission"))
    add(("subtitle", f"Response to the Client's letter and Attachments A to C of 21 October 2026  |  Release {a.release}  |  {a.date}"))

    add(("h1", "1. Submission notice (Schedule B1)"))
    add(("table", [["Item", "Detail"],
                   ["Milestone", "2: First integrated build (Schedule B). Also the Schedule A6 early technical proof (Delivery Plan)."],
                   ["Release identifier", a.release],
                   ["Repository commit", f"{a.commit} (git tag {a.release})"],
                   ["Staging website", a.web],
                   ["Staging API", f"{a.api}  (GET /api/version reports the commit above)"],
                   ["Database", "Migrations 001 to 016, recorded with their SHA-256 in dbo.SchemaMigration. 016 is new in this release and additive."],
                   ["Evidence pack", f"PHILMART_M2_Evidence_{a.release}.zip, in the structure proposed in Attachment C"],
                   ["Prepared by", "Ratnesh (developer)"],
                   ["Declared by", "Morne (lead developer), section 10"]], [4.2, 12.8]))
    add(("p", "Under clause 10, the review period begins once this package is materially complete. We believe it now is: "
              "every item in Schedule B1 is included, and every one of the 130 tests has a recorded outcome."))

    add(("h1", "2. Our response to your proposal"))
    add(("p", "Thank you for the careful review and for proposing a practical method. We agree with it. "
              "In particular, we agree that evidence produced once and reviewed independently is better for both parties than repeating the same tests. "
              "This submission uses Milestone 2 as the pilot you suggested."))
    add(("numbered", [
        "<b>The 30 additional tests are incorporated.</b> We merged N01 to N30 into the 100-row checklist. Where two tests cover the same thing we say so, and we did not remove any test. All 130 now map to the D076 requirement they verify (section 5 and the index).",
        "<b>Existing evidence was reused, and everything was run again.</b> The automated suite was re-run in full against this release. Every checklist item was then re-executed on " + env_label + " by an automated browser and API run. That run saved screenshots, sanitised API exchanges and database records for each test ID.",
        "<b>The evidence pack follows Attachment C</b>, including the index (00_INDEX.xlsx), release record, machine-readable test output, screenshots, API and database traces, security and auction folders, defects and run instructions. The scripts that produced it are in the repository (tools/evidence), so you or your reviewer can run them again.",
        "<b>Outcomes are reported strictly.</b> PASS means a test was executed and passed. Anything we could not execute on that environment is BLOCKED, with the reason. Nothing unexecuted is reported as PASS.",
        "<b>New gaps your tests found have been fixed, not only evidenced.</b> Sign-out did not end the session on the server (N03). The PIN limit was per address only, not per client (N01). There was no way to identify the running build (N28). The A6 invoice, payment and reversal proof (N29) had no screen. All four are now in this release.",
    ]))
    add(("p", "We have four requests, so that the method also works well for Milestones 3 and 4:"))
    add(("numbered", [
        "<b>Review period.</b> We understand the 15 working days in clause 10 run from receipt of this package, and that the decision is made against the agreed D076 requirements. Isolated cosmetic or Low Severity items are recorded for correction rather than holding acceptance (clause 10, final paragraph).",
        "<b>The Requirements and Testing Matrix.</b> We welcome it and will review it promptly. We ask that it is agreed before the Milestone 3 submission and then held stable for that milestone. Any test added later can be scheduled for the next milestone or agreed as a change, so each submission is measured against a known list.",
        "<b>Milestone boundaries.</b> Some proposed tests belong to later milestones, as Attachment C section 6 itself notes. N30 (backup and restore) is part of the Schedule C handover; we demonstrated it now on a copy (section 5) and propose the formal hosted restore at the Milestone 4 continuity test. N24 is complete for the emails Milestone 2 sends; the full notification layer is task T25 (Milestone 4).",
        "<b>A separate staging site.</b> philmart.co.za currently serves the same site and database as staging. Before any production release we recommend separating them, so that tests never run on production data. This needs a decision on hosting, which is yours.",
    ]))

    add(("h1", "3. Results at a glance"))
    tbl = [["Source", "PASS", "FAIL", "BLOCKED", "NOT RUN", "Clarification / later milestone", "Total"]]
    for src, label in (("Checklist 19 Oct", "Original checklist (100)"), ("Attachment A", "Additional tests N01 to N30"), ("Attachment B", "Interim findings B01 to B14")):
        other = total(src) - sum(c(src, k) for k in ("PASS", "FAIL", "BLOCKED", "NOT RUN"))
        tbl.append([label, str(c(src, "PASS")), str(c(src, "FAIL")), str(c(src, "BLOCKED")), str(c(src, "NOT RUN")), str(other), str(total(src))])
    add(("table", tbl, [6, 1.6, 1.6, 2, 2, 3, 1.4]))
    add(("p", f"<b>Automated tests:</b> {auto['passed']} of {auto['total']} passed (unit and integration, against a database built from migrations 001 to 016)."))
    if env:
        add(("p", f"<b>Environment run</b> ({env['label']}, {env['web']}, run {env['runId']}): {env['pass']} passed, {env['fail']} failed, {env['blocked']} blocked, out of {env['total']} checks."))
    blocked_ids = sorted([k for k, v in rows.items() if v["outcome"] == "BLOCKED"])
    failed_ids = sorted([k for k, v in rows.items() if v["outcome"] == "FAIL"])
    if failed_ids:
        add(("p", "<b>Failed:</b> " + ", ".join(failed_ids) + ". See the index and 07_DEFECTS."))
    else:
        add(("p", "<b>No test failed.</b>"))
    if blocked_ids:
        add(("p", "<b>Blocked:</b> " + ", ".join(blocked_ids) + ". The reasons are in section 9."))

    add(("h1", "4. Item-by-item response to Attachment B"))
    add(("p", "Your findings were accurate and helpful. B01 to B05 and B08 were real and are fixed. B06, B07, B09, B10 and B11 are clarifications or Client dependencies, with our proposed position. For B12 to B14, the evidence is now supplied."))
    b_resp = {
        "B01": ("Confirmed defect. The price sort used an auction's starting price, while the card shows the current bid. An auction that started at R5 and stood at R600 was therefore sorted as R5.",
                "Fixed. Both price sorts now use the price shown on the card: the fixed price, otherwise the current bid, otherwise the starting price. A test reproduces your R600 / R5 / R42 case."),
        "B02": ("Confirmed on screen. The server already refused a blank street or city, so no sale could be created with that address. The confirm button should not have been enabled, though.",
                "Fixed. Confirm Purchase stays disabled until street, city and postal code are filled in, and the destination line no longer reads 'Destination: ,'. The server now also requires the postal code."),
        "B03": ("Confirmed on screen. The server has always refused a bid below the minimum (database rule P_List_Bid_Place, error 50036), and no such bid can be recorded. The button should not have been enabled.",
                "Fixed. The amount turns red with 'Your bid must be at least R ...' and Confirm Bid stays disabled. A new test also proves that a refused bid leaves no commitment record behind (N21)."),
        "B04": ("Confirmed. A live auction refreshes every 10 seconds, and each refresh rebuilt the panel, which closed the history.",
                "Fixed. The history stays open across refreshes. The screenshot was taken 14 seconds after opening it."),
        "B05": ("Agreed. The line appeared even where nothing could be changed.",
                "Fixed. It now appears only where an address or pickup point can be changed for this purchase. It is not shown for Collect from the Shop or for bids."),
        "B06": ("Clarification. D076 s4.3 and s22 allow an override 'where the purchase flow permits'. A bid does not create the transaction: the sale and its delivery snapshot are created when the auction is won. In our design, delivery for an auction win is confirmed when the Shop prepares the invoice (BR-15, Milestone 3), and the winner can choose the method and address at that point.",
                "The bid screen now says this in plain words. Please confirm this reading of D076. If you would prefer an address override at bid time as well, we will add it, with no database change needed."),
        "B07": ("Agreed: a release dependency, not a defect. Agreement A5 asks us to build the placement, versioning and acceptance mechanics, but not to publish draft wording as approved.",
                "The mechanics are complete and tested (N04, N09). The approved text needs no code change: a new published version is loaded, and buyers are asked to accept it before their next purchase or bid. We need the approved text from you before production."),
        "B08": ("Agreed. With a search active, the counter counted matches but read as if nothing was saved.",
                "Fixed. It now reads '0 saved items match \"...\"' and offers 'Show all saved items'."),
        "B09": ("The storefront shows the public profile name and blurb when a Shop has them. The logo, profile and contact details are captured in Shop Setup (BR-07, Milestone 3, task T17), and the demo Shops have none entered.",
                "No change for Milestone 2. The full storefront profile is delivered with Shop Setup in Milestone 3."),
        "B10": ("Intended temporary state. The Shop application flow is BR-06 (Milestone 3, task T16). Until then, the page explains the process and opens Support with 'Opening a Shop' selected.",
                "No change. Please confirm this is acceptable for Milestone 2. It is replaced by the application flow in Milestone 3."),
        "B11": ("Agreed: no filter defect. The demo items carry sample classifications only (D053: sample values are not rules). The region filter includes the countries under a region, and a test covers this.",
                "No taxonomy change made, as none is authorised. We will apply the outcome of your taxonomy review."),
        "B12": ("Agreed: an evidence gap.",
                "Evidence supplied: automated tests for simultaneous bids, ordering, the exact close boundary, repeated soft-close extensions, simultaneous and repeated closes, and outage recovery. Also the environment run (A.10 to A.13, O.2, O.3, W.9, W.10) and a backup, restore and restart demonstration (O.1, N23, N30)."),
        "B13": ("Agreed: an evidence gap.",
                "Evidence supplied: 8 isolation tests plus N16, N18 and N19. The environment run includes cross-Shop API attempts (I.1 to I.5) and a staff session losing access the moment permissions change (S.6 to S.11). The API requests and denial responses are saved in 05_SECURITY."),
        "B14": ("Agreed. In the 19 October checklist, B.4, S.5 and S.6 were F and R.10, B.10, S.8 and G.5 were blank. The fixes went out on 19 October, but the build was not tied to a repository commit.",
                f"Every row has now been re-executed against tagged commit {a.commit[:12]} ({a.release}). The API reports its commit at /api/version. The index maps each test to its requirement and evidence."),
    }
    tb = [["Ref", "Finding", "Our assessment", "Action in this release", "Outcome"]]
    for (tid, title, *_rest) in ATTACHMENT_B:
        assess, action = b_resp[tid]
        tb.append([tid, title, assess, action, outcome(tid)])
    add(("table", tb, [1.1, 2.8, 6.0, 6.0, 1.6], "small"))

    add(("h1", "5. Additional tests N01 to N30"))
    add(("p", "Each test below has an automated test, an environment check, or both. 'Overlaps' names the original checklist rows that cover the same ground. We kept both, as both add evidence."))
    ta = [["ID", "Test", "Outcome", "How it is evidenced", "Overlaps"]]
    for (tid, title, expected, req, milestone, names, overlap, note) in ATTACHMENT_A:
        how = (f"{len(names)} automated test{'s' if len(names) != 1 else ''}" if names else "Environment run") + (". " + note if note else "")
        ta.append([tid, title, outcome(tid), how, overlap])
    add(("table", ta, [1.1, 4.2, 1.6, 8.0, 2.6], "small"))
    add(("p", "<b>N29 and the Schedule A6 technical proof.</b> The Delivery Plan says Milestone 2 doubles as the A6 early technical proof. That proof includes an invoice, a Shop-recorded buyer payment and a permitted reversal or correction. These now work end to end: in the Shop workspace under Sales &amp; Invoices, and in the API. They rely on the database rules already in place (D019, an issued invoice cannot be edited; D021, payment allocation; D022, reversals are separate linked records). The full invoicing and statement screens (BR-15, BR-16) remain in Milestones 3 and 4, as planned."))
    add(("p", "<b>N30 backup, restore and restart.</b> This was demonstrated on a copy of the database: a COPY_ONLY backup, a restore under a new name, migrations and row counts compared, auctions left overdue while the API was down, then the API restarted on the restored copy. Its start-up sweep closed every overdue auction correctly, with a winner, a sale and an email, or unsold with the item returned to stock. We did not restore the hosted database, because that needs the hosting panel and an isolated copy (Attachment C, section 5). We propose to do it at the Milestone 4 continuity test."))

    add(("h1", "6. Release notes"))
    add(("p", f"<b>Changes in {a.release}</b> since the build you tested on 19 October:"))
    add(("bullets", [
        "B01: price sorting uses the current bid of an auction.",
        "B02: an incomplete alternative address can't be confirmed; postal code required on the server too.",
        "B03: a bid below the minimum can't be confirmed on screen (the server already refused it).",
        "B04: Bid History stays open across the live refresh.",
        "B05: 'a change here applies to this purchase only' shown only where something can change; clear wording on the bid screen (B06).",
        "B08: saved-items counter wording during a search, with 'Show all saved items'.",
        "N01: a per-client limit on PIN requests (30 per 10 minutes), on top of the per-address limit.",
        "N03: Sign out now ends the session on the server. The token is recorded as revoked and refused from then on (migration 016).",
        "N28: GET /api/version reports the version and source commit of the running API.",
        "N29 / Schedule A6: Sales & Invoices in the Shop workspace: issue an invoice for a sale, record a buyer payment, reverse it with a reason. Each action needs its own permission.",
        "N02 / L.5: a resent PIN now cancels the earlier one at the API too, not only on screen.",
        "G.1 (found by our own evidence run): on a 375 px phone the marketplace scrolled sideways by 27 px because of the page-number row. It now wraps. Prices no longer break between 'R' and the amount.",
        "28 new automated tests, each tagged with its review ID (e.g. dotnet test --filter Review=N03).",
        "tools/evidence: the evidence runner, the outage and restore script and the index builder used for this pack.",
    ]))
    add(("p", "<b>What Milestone 2 implements</b> (Schedule B trigger): buyer registration and login with versioned legal acceptance (BR-01); the public marketplace, storefronts, Saved Items, Recently Viewed and support (BR-02); item detail, Buy Now and Confirm Bid with immutable delivery snapshots (BR-03); Shop users, direct permissions and Shop isolation enforced in the API and the database (BR-08); auctions with authoritative server time, soft close, safe concurrency, closing, and outage recovery (BR-13). Also emails EML-012, -013, -015, -019 and -020; buyer restrictions (D044); Shop deactivation by PHILMART (D031); and now the A6 invoice, payment and reversal proof. Requirement-by-requirement status is in docs/milestone-2/README.md in the repository."))

    add(("h1", "7. Known issues and declaration (Schedule B1)"))
    add(("p", "There are <b>no known Critical or High Severity defects</b>. These are the known issues, including every one we consider security-relevant:"))
    add(("table", [
        ["Issue", "Severity", "Security", "Plan"],
        ["Legal documents and commitment wording are placeholders (version 0.1-dev). B07.", "Release dependency", "No", "Client to supply the approved text. Load it as a new version before production."],
        ["Emails from philmart.co.za can land in spam or be rejected. The domain publishes SPF -all and DMARC p=reject but has no DKIM signing, and the mail server is not in the SPF record.", "Medium", "Yes (spoofing / delivery)", "Enable DKIM in the hosting mail panel, or relax DMARC to p=none until it is. A DNS and hosting action, not code."],
        ["Found by the live evidence run: the hosting mail server limits how fast PHILMART may send. After about 40 emails in a short time it refused further ones ('Sender is throttled'), and those were marked failed after their retries. A busy auction day could reach the same limit.", "Medium", "No", "Ask the host to raise the sending limit, or use a transactional email service. Re-queue messages refused by a throttle once it clears: email layer T25 (Milestone 4), sooner if you prefer."],
        ["The browser keeps the sign-in token in local storage, so a script-injection flaw could expose it. Sign out now revokes it on the server (N03), and the site loads no third-party scripts.", "Medium", "Yes", "Move to an httpOnly cookie once the site and API share one domain (a hosting decision). Planned for the Milestone 4 security review (T37)."],
        ["Two-factor sign-in for PHILMART administrators is not built yet.", "Medium", "Yes", "BR-21, Milestone 4, as noted in the checklist."],
        ["Wrong-PIN counters and per-client limits are held in the API's memory. This is correct on one server, but a second server would need a shared cache.", "Low", "Yes", "Fine on the current single server. Revisit if the hosting is scaled out."],
        ["philmart.co.za and staging are the same site and database.", "Low", "Yes (data separation)", "Separate them before production (section 2, request 4)."],
        ["Found by the live evidence run: when a buyer resumes registration, the page shows an empty Personal Details step for a second or two before moving to the step they reached (R.7). Resuming itself works.", "Low", "No", "Fixed in the repository (commit 8671681): it now shows 'Loading your registration'. Goes out with the next upload."],
        ["Delivery tariffs are not loaded, so invoices carry no shipping line yet.", "Low", "No", "Delivery tariffs and configuration: Milestones 3 and 4 (BR-15, BR-22)."],
        ["Not in Milestone 2: Ask Shop a Question (M4), auction events (M4), bulk upload (M3), item images (M3), full invoicing and statements (M3/M4).", "n/a", "No", "As in the Delivery Plan."],
    ], [7.0, 2.0, 2.2, 5.8], "small"))

    add(("h1", "8. Configuration and migration notes"))
    add(("bullets", [
        "Migration 016_revoked_session.sql adds one table (philmart.Sys_RevokedSession). It is additive, and the API from 19 October keeps working if it is applied first.",
        "New optional setting Auth:PinRequestsPer10Minutes (default 30). No other settings changed. The server's own appsettings (mail, connection string, CORS) are not part of the upload.",
        "Deployment order: database, then API, then website. The steps are in UPLOAD_STEPS.txt in the deployment package. After deployment, /api/version must report commit " + a.commit[:12] + ".",
    ]))

    add(("h1", "9. Limits of this evidence"))
    add(("p", "To be clear about what the evidence does and does not show:"))
    add(("bullets", [
        "<b>Inbox arrival.</b> The run proves each email was produced for the right person, with the right content, and handed to the mail server. Two messages are never stored, by design: the sign-in PIN and the support copy. For these, and for arrival in a real inbox (L.3, L.5, E.3, E.8, E.9), the mailbox is needed. These rows were marked Pass in the signed-off checklist of 19 October and in your own client-assisted testing, and the automated tests prove the logic. On live, the hosting mail server also refused a few emails near the end of the run because it limits how fast we may send (section 7); those rows are BLOCKED with the mail server's reply.",
        "<b>Outage on the hosted server.</b> Restarting the hosted API needs the hosting panel. O.1, N23 and N30 were demonstrated on a restored copy with a real stop and restart of the API process. The same recovery code runs every 5 seconds on the hosted server (O.2).",
        "<b>Browsers.</b> Chrome and Edge were automated, plus Firefox and Safari engines where installed. Your iPhone Safari testing remains the reference for iOS.",
        "<b>A first live attempt.</b> The first live run scheduled its test auctions at its very start. On the live site the earlier sections took longer than those auctions ran, so they had closed, correctly and on time, before the bidding checks began, and those checks failed for that reason. The runner now schedules the auctions just before the auction checks, and the run was repeated in full. The first attempt is kept unedited in 02_TEST_OUTPUT/earlier-live-attempt; the index uses the complete second run.",
                "<b>Where it ran, and what it changed.</b> Your execution rules (Attachments A and C) ask for purchases, bids, permission changes, API security probes and concurrent writes only with prior approval and in isolated staging. Today the staging site is also philmart.co.za (section 2, request 4), so the run worked on the live database, using only synthetic buyers (addresses starting m2-, on a non-deliverable test domain), the two demo Shops and demo items. It changed: a fresh batch of demo listings and a few Ready to List demo items; Buy Now purchases and bids by the synthetic buyers; test auctions on Cape Stamp Traders, scheduled, cancelled and closed; one synthetic Shop user invited; staff@capestamps.test's permissions changed and restored; one synthetic buyer restricted and lifted at Highveld Philatelics; Highveld Philatelics deactivated for under a minute (M.6) and reactivated; one support request marked 'M2 evidence test - please ignore'; and one invoice, payment and reversal on a demo sale. Every change is in the audit log. No real money moved and no real person was contacted. The direct API probes only tried to cross between the two demo Shops, and all were refused. If you would prefer these checks to run only on a separate staging copy in future, we will set one up.",
    ]))

    add(("h1", "10. Developer declaration"))
    add(("p", f"I have checked this submission for completeness. It relates to release {a.release} (commit {a.commit}). The test results are those produced by the automated tests and the evidence run described above, and none was edited by hand. Every known Critical, High or security-relevant issue is declared in section 7. AI tools assisted in writing code, tests and this document; the evidence itself is machine-generated test output, screenshots and database records, and all of it was reviewed by the developers. No personal data, passwords, PINs or tokens are included in the evidence pack."))
    add(("signatures", [("Morne", "Lead developer"), ("Ratnesh", "Developer (prepared this submission)")]))
    return B


# --------------------------------------------------------------- DOCX

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def runs_from(par, text, size=None, bold_all=False, color=None):
    # tiny <b>...</b> support
    parts = text.replace("&amp;", "&").split("<b>")
    first = True
    for chunk in parts:
        if first:
            seg = [(chunk, False)]
            first = False
        else:
            b, _, rest = chunk.partition("</b>")
            seg = [(b, True), (rest, False)]
        for t, bold in seg:
            if not t:
                continue
            r = par.add_run(t)
            r.bold = bold or bold_all
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)


def write_docx(blocks, path, header):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    hp = sec.header.paragraphs[0]
    runs_from(hp, header, size=8, color="667085")
    fp = sec.footer.paragraphs[0]
    runs_from(fp, header, size=8, color="667085")

    for kind, payload, *rest in blocks:
        if kind == "title":
            p = doc.add_paragraph()
            runs_from(p, "MILESTONE 2", size=10, bold_all=True, color=BLUE)
            p = doc.add_paragraph()
            runs_from(p, payload, size=20, bold_all=True, color=NAVY)
        elif kind == "subtitle":
            p = doc.add_paragraph()
            runs_from(p, payload, size=9, color="475467")
        elif kind == "h1":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            runs_from(p, payload, size=13, bold_all=True, color=BLUE)
        elif kind == "p":
            p = doc.add_paragraph()
            runs_from(p, payload)
        elif kind in ("bullets", "numbered"):
            for item in payload:
                p = doc.add_paragraph(style="List Bullet" if kind == "bullets" else "List Number")
                runs_from(p, item)
        elif kind == "table":
            widths = rest[0]
            small = len(rest) > 1 and rest[1] == "small"
            t = doc.add_table(rows=0, cols=len(payload[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, row in enumerate(payload):
                cells = t.add_row().cells
                for j, val in enumerate(row):
                    cells[j].width = Cm(widths[j])
                    par = cells[j].paragraphs[0]
                    runs_from(par, str(val), size=8 if small else 9, bold_all=(i == 0), color="FFFFFF" if i == 0 else None)
                    if i == 0:
                        shade(cells[j], NAVY)
                    elif val in ("PASS",):
                        shade(cells[j], "D9F2E3")
                    elif val in ("FAIL",):
                        shade(cells[j], "F8D7DA")
                    elif val in ("BLOCKED",):
                        shade(cells[j], "FFF3CD")
            doc.add_paragraph()
        elif kind == "signatures":
            t = doc.add_table(rows=4, cols=len(payload))
            for j, (name, role) in enumerate(payload):
                t.cell(0, j).paragraphs[0].add_run("Signature: ______________________________")
                t.cell(1, j).paragraphs[0].add_run(f"Name: {name}")
                t.cell(2, j).paragraphs[0].add_run(f"Role: {role}")
                t.cell(3, j).paragraphs[0].add_run("Date: ____________________")
    doc.save(path)


# --------------------------------------------------------------- HTML

CSS = """
@page { size: A4; margin: 18mm 16mm 18mm 16mm; }
body { font-family: 'DejaVu Sans', 'Segoe UI', Arial, sans-serif; font-size: 9.6pt; color: #1d2939; line-height: 1.42; }
.kicker { color: #2F6DA8; font-weight: 700; font-size: 10pt; letter-spacing: .02em; margin: 0 0 2mm; }
h1.title { color: #16233A; font-size: 21pt; margin: 0 0 2mm; line-height: 1.15; }
.subtitle { color: #475467; font-size: 9pt; margin: 0 0 6mm; }
h2 { color: #2F6DA8; font-size: 12.5pt; margin: 7mm 0 2.5mm; break-after: avoid; }
p { margin: 0 0 2.6mm; }
ul, ol { margin: 0 0 3mm 5mm; padding-left: 4mm; }
li { margin: 0 0 1.6mm; }
table { width: 100%; border-collapse: collapse; margin: 1mm 0 4mm; table-layout: fixed; }
th { background: #16233A; color: #fff; text-align: left; font-weight: 700; padding: 2mm; font-size: 8.6pt; vertical-align: top; }
td { border: 0.3mm solid #d0d5dd; padding: 1.8mm 2mm; vertical-align: top; font-size: 8.8pt; }
table.small td, table.small th { font-size: 7.9pt; padding: 1.5mm 1.7mm; }
tr:nth-child(even) td { background: #f5f7fa; }
td.PASS { background: #d9f2e3 !important; font-weight: 700; }
td.FAIL { background: #f8d7da !important; font-weight: 700; }
td.BLOCKED { background: #fff3cd !important; font-weight: 700; }
tr { break-inside: avoid; }
.sig { display: grid; grid-template-columns: 1fr 1fr; gap: 10mm; margin-top: 8mm; }
.sig div p { margin: 0 0 5mm; }
"""


def inline(text):
    # keep <b> from the content, escape everything else
    esc = html.escape(text.replace("&amp;", "&"), quote=False)
    return esc.replace("&lt;b&gt;", "<b>").replace("&lt;/b&gt;", "</b>")


def write_html(blocks, path, header):
    out = [f"<!doctype html><html><head><meta charset='utf-8'><title>PHILMART M2 submission</title><style>{CSS}</style></head><body>"]
    for kind, payload, *rest in blocks:
        if kind == "title":
            out.append("<p class='kicker'>MILESTONE 2</p>")
            out.append(f"<h1 class='title'>{inline(payload)}</h1>")
        elif kind == "subtitle":
            out.append(f"<p class='subtitle'>{inline(payload)}</p>")
        elif kind == "h1":
            out.append(f"<h2>{inline(payload)}</h2>")
        elif kind == "p":
            out.append(f"<p>{inline(payload)}</p>")
        elif kind in ("bullets", "numbered"):
            tag = "ul" if kind == "bullets" else "ol"
            out.append(f"<{tag}>" + "".join(f"<li>{inline(i)}</li>" for i in payload) + f"</{tag}>")
        elif kind == "table":
            widths = rest[0]
            cls = "small" if len(rest) > 1 and rest[1] == "small" else ""
            total = sum(widths)
            cols = "".join(f"<col style='width:{w / total * 100:.1f}%'>" for w in widths)
            head = "".join(f"<th>{inline(str(x))}</th>" for x in payload[0])
            body = "".join("<tr>" + "".join(f"<td class='{v if v in ('PASS', 'FAIL', 'BLOCKED') else ''}'>{inline(str(v))}</td>" for v in row) + "</tr>" for row in payload[1:])
            out.append(f"<table class='{cls}'><colgroup>{cols}</colgroup><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
        elif kind == "signatures":
            out.append("<div class='sig'>" + "".join(
                f"<div><p>Signature: ______________________________</p><p>Name: {html.escape(n)}</p><p>Role: {html.escape(r)}</p><p>Date: ____________________</p></div>"
                for n, r in payload) + "</div>")
    out.append("</body></html>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("".join(out))
    with open(path + ".header.txt", "w", encoding="utf-8") as f:
        f.write(header)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--facts", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--release", required=True)
    ap.add_argument("--commit", required=True)
    ap.add_argument("--web", required=True)
    ap.add_argument("--api", required=True)
    ap.add_argument("--date", required=True)
    a = ap.parse_args()
    with open(a.facts, encoding="utf-8") as fh:
        facts = json.load(fh)
    os.makedirs(a.out, exist_ok=True)
    blocks = content(facts, a)
    header = f"PHILMART | Milestone 2 developer response and evidence submission | {a.release} | {a.date}"
    base = os.path.join(a.out, f"PHILMART_M2_Developer_Response_and_Evidence_{a.release}")
    write_docx(blocks, base + ".docx", header)
    write_html(blocks, base + ".html", header)
    print(base + ".docx")
    print(base + ".html")


if __name__ == "__main__":
    main()
